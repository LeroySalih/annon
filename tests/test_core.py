import io
import zipfile

import docx
import pymupdf

from annon import core

ROSTER = b"name\nAda Lovelace\nAlan Turing\nGrace Hopper\n"


def _pdf():
    d = pymupdf.open()
    for _ in range(2):
        p = d.new_page()
        p.insert_text((72, 300), "Question 1")
        p.insert_text((200, p.rect.height - 25), "OLD FOOTER Page 1 of 6")
    return d.tobytes()


def _docx():
    d = docx.Document()
    d.add_paragraph("Question 1")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def test_roster_variants():
    assert core.read_roster(b"First,Last\nA,B\n") == ["A B"]
    assert core.read_roster(b"x\nA\n") == ["A"]


def test_codes_unique_5_digit():
    pairs = core.assign_codes([f"p{i}" for i in range(2000)])
    codes = [c for _, c in pairs]
    assert len(set(codes)) == 2000 and all(len(c) == 5 for c in codes)


def test_personalise_pdf_header_and_footer():
    z = zipfile.ZipFile(io.BytesIO(core.personalise(_pdf(), "pdf", ROSTER, title="Systems")))
    key = dict(l.split(",")[::-1] for l in z.read("codes.csv").decode().split("\r\n")[1:] if l)
    code = key["Ada Lovelace"]
    doc = pymupdf.open(stream=z.read(f"{code}.pdf"), filetype="pdf")
    for page in doc:
        h = page.rect.height
        top = page.get_text(clip=pymupdf.Rect(0, 0, page.rect.width, 50))
        bot = page.get_text(clip=pymupdf.Rect(0, h - 50, page.rect.width, h))
        assert "Ada Lovelace" in top and code in top
        assert "OLD FOOTER" not in bot and "Ada" not in bot
        assert code in bot and "Systems" in bot and f"Page {page.number + 1} of 2" in bot
        # code left, title centre, page right
        w = {x[4]: x[0] for x in page.get_text("words", clip=pymupdf.Rect(0, h - 50, page.rect.width, h))}
        assert w[code] < w["Systems"] < w["Page"]


def test_personalise_docx():
    z = zipfile.ZipFile(io.BytesIO(core.personalise(_docx(), "docx", ROSTER, title="Systems")))
    key = dict(l.split(",")[::-1] for l in z.read("codes.csv").decode().split("\r\n")[1:] if l)
    code = key["Alan Turing"]
    n = f"{code}.docx"
    assert n in z.namelist()
    d = docx.Document(io.BytesIO(z.read(n)))
    s = d.sections[0]
    assert "Alan Turing" in s.header.paragraphs[-1].text and code in s.header.paragraphs[-1].text
    ft = s.footer.paragraphs[-1]
    assert len(s.footer.paragraphs) == 1 and ft.text.startswith(f"Code: {code}")
    assert "Systems" in ft.text and "PAGE" in ft._p.xml and "NUMPAGES" in ft._p.xml


def test_redact_blocks_header_keeps_footer():
    src = core.personalise_pdf(_pdf(), "Ada Lovelace", "12345")
    out = pymupdf.open(stream=core.redact_scan(src, "pdf", band_mm=25), filetype="pdf")
    assert len(out) == 2
    pix = out[0].get_pixmap(dpi=72)
    assert pix.pixel(10, 10)[:3] == (0, 0, 0)
    assert out[0].get_text() == ""  # flattened: no recoverable text


def test_app_endpoints():
    from annon.app import app
    c = app.test_client()
    r = c.post("/personalise", data={"assignment": (io.BytesIO(_pdf()), "a.pdf"), "roster": (io.BytesIO(ROSTER), "r.csv")})
    assert r.status_code == 200 and r.mimetype == "application/zip"
    r = c.post("/redact", data={"scan": (io.BytesIO(_pdf()), "s.pdf")})
    assert r.status_code == 200
    assert c.post("/redact").status_code == 400
