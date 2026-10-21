"""Personalise PDF/DOCX assignments per pupil and redact headers on scans.

Pupil names stay in memory/output files only; nothing here logs them.
"""
import csv
import io
import os
import pathlib
import re
import secrets
import shutil
import subprocess
import tempfile

import docx
import pymupdf

MM = 72 / 25.4  # PDF points per millimetre


def read_roster(data: bytes) -> list[str]:
    """Return pupil names from CSV bytes.

    Uses a 'name' column, else 'first'+'last' columns, else the first column.
    """
    text = data.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        raise ValueError("Class list is empty")
    head = [c.strip().lower() for c in rows[0]]
    body = rows[1:]
    if "name" in head:
        i = head.index("name")
        names = [r[i] for r in body if len(r) > i]
    else:
        first = next((i for i, h in enumerate(head) if h in ("first", "first name", "firstname", "forename")), None)
        last = next((i for i, h in enumerate(head) if h in ("last", "last name", "lastname", "surname")), None)
        if first is not None and last is not None:
            names = [f"{r[first]} {r[last]}" for r in body if len(r) > max(first, last)]
        else:
            names = [r[0] for r in body]
    names = [re.sub(r"\s+", " ", n).strip() for n in names]
    names = [n for n in names if n]
    if not names:
        raise ValueError("No pupil names found")
    return names


def assign_codes(names: list[str]) -> list[tuple[str, str]]:
    """Give each pupil a unique random 5-digit code (10000-99999)."""
    if len(names) > 90000:
        raise ValueError("Too many pupils for 5-digit codes")
    used: set[str] = set()
    out = []
    for n in names:
        while True:
            c = str(10000 + secrets.randbelow(90000))
            if c not in used:
                used.add(c)
                break
        out.append((n, c))
    return out


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_") or "pupil"


def personalise_pdf(src: bytes, name: str, code: str, margin_mm: float = 12,
                    title: str = "", footer_mm: float = 20) -> bytes:
    """Name+code in the top margin; existing footer text replaced by
    code (left), title (centre), 'Page X of Y' (right)."""
    doc = pymupdf.open(stream=src, filetype="pdf")
    y = margin_mm * MM / 2 + 3
    pad = 36
    total = len(doc)
    for page in doc:
        r = page.rect
        page.add_redact_annot(pymupdf.Rect(0, r.height - footer_mm * MM, r.width, r.height), fill=(1, 1, 1))
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
        page.insert_text((pad, y), name, fontsize=11, fontname="helv")
        label = f"Pupil ID: {code}"
        w = pymupdf.get_text_length(label, fontname="helv", fontsize=11)
        page.insert_text((r.width - pad - w, y), label, fontsize=11, fontname="helv")
        fy = r.height - y + 6
        page.insert_text((pad, fy), label, fontsize=10, fontname="helv")
        if title:
            tw = pymupdf.get_text_length(title, fontname="helv", fontsize=10)
            page.insert_text(((r.width - tw) / 2, fy), title, fontsize=10, fontname="helv")
        pg = f"Page {page.number + 1} of {total}"
        pw = pymupdf.get_text_length(pg, fontname="helv", fontsize=10)
        page.insert_text((r.width - pad - pw, fy), pg, fontsize=10, fontname="helv")
    out = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    return out


def _field(par, instr: str):
    """Append a Word field (e.g. PAGE, NUMPAGES) to a paragraph."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    f = OxmlElement("w:fldSimple")
    f.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    f.append(r)
    par._p.append(f)


def _clear(part):
    """Remove all paragraphs/tables from a header/footer, leaving one empty paragraph."""
    for el in list(part._element):
        if el.tag.endswith("}p") or el.tag.endswith("}tbl"):
            part._element.remove(el)
    return part.add_paragraph()


def personalise_docx(src: bytes, name: str, code: str, title: str = "") -> bytes:
    from docx.enum.text import WD_TAB_ALIGNMENT
    d = docx.Document(io.BytesIO(src))
    for s in d.sections:
        heads = [s.header]
        foots = [s.footer]
        if s.different_first_page_header_footer:
            heads.append(s.first_page_header)
            foots.append(s.first_page_footer)
        width = s.page_width - s.left_margin - s.right_margin
        for h in heads:
            h.is_linked_to_previous = False
            p = h.add_paragraph(f"{name}\tPupil ID: {code}")
            p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        for f in foots:
            f.is_linked_to_previous = False
            p = _clear(f)
            ts = p.paragraph_format.tab_stops
            ts.add_tab_stop(width // 2, WD_TAB_ALIGNMENT.CENTER)
            ts.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
            p.add_run(f"Pupil ID: {code}\t{title}\tPage ")
            _field(p, "PAGE")
            p.add_run(" of ")
            _field(p, "NUMPAGES")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _soffice() -> str:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    for p in ("/Applications/LibreOffice.app/Contents/MacOS/soffice",
              r"C:\\Program Files\\LibreOffice\\program\\soffice.exe"):
        if not exe and os.path.exists(p):
            exe = p
    if not exe:
        raise ValueError("DOCX needs LibreOffice (or Microsoft Word on Windows) to make the combined PDF")
    return exe


_WORD_PS = r"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {
  foreach ($f in Get-Content -LiteralPath $args[0]) {
    $doc = $word.Documents.Open($f, $false, $true)
    $doc.ExportAsFixedFormat(($f -replace '\.docx$', '.pdf'), 17)
    $doc.Close($false)
  }
} finally { $word.Quit() }
"""


def _word_to_pdfs(tmp: str, n: int) -> None:
    """Convert tmp/0.docx.. with Microsoft Word (Windows, no admin rights needed)."""
    listing = os.path.join(tmp, "files.txt")
    with open(listing, "w", encoding="utf-8") as f:
        f.write("\n".join(os.path.join(tmp, f"{i}.docx") for i in range(n)))
    script = os.path.join(tmp, "word.ps1")
    with open(script, "w", encoding="utf-8") as f:
        f.write(_WORD_PS)
    r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", script, listing],
                       capture_output=True, text=True, timeout=600)
    if r.returncode != 0:
        raise ValueError("Microsoft Word could not convert the DOCX: " + r.stderr.strip()[:300])


def docx_to_pdfs(docs: list[bytes]) -> list[bytes]:
    """Convert DOCX files to PDF, keeping order. Uses LibreOffice if found,
    else Microsoft Word on Windows."""
    try:
        exe = _soffice()
    except ValueError:
        if os.name != "nt":
            raise
        exe = None
    with tempfile.TemporaryDirectory() as tmp:
        for i, data in enumerate(docs):
            with open(os.path.join(tmp, f"{i}.docx"), "wb") as f:
                f.write(data)
        if exe:
            cmd = [exe, f"-env:UserInstallation=file:///{pathlib.Path(tmp).as_posix().lstrip('/')}/profile",
                   "--headless", "--convert-to", "pdf", "--outdir", tmp] + \
                  [os.path.join(tmp, f"{i}.docx") for i in range(len(docs))]
            subprocess.run(cmd, check=True, capture_output=True, timeout=300)
        else:
            _word_to_pdfs(tmp, len(docs))
        out = []
        for i in range(len(docs)):
            with open(os.path.join(tmp, f"{i}.pdf"), "rb") as f:
                out.append(f.read())
    return out


def personalise(src: bytes, ext: str, roster: bytes, margin_mm: float = 12,
                title: str = "", footer_mm: float = 20) -> tuple[bytes, str]:
    """Return (one PDF with a personalised copy per pupil in roster order,
    CSV text mapping pupil ID -> name)."""
    ext = ext.lower().lstrip(".")
    if ext not in ("pdf", "docx"):
        raise ValueError("Assignment must be a .pdf or .docx")
    pairs = assign_codes(read_roster(roster))
    if ext == "pdf":
        pdfs = [personalise_pdf(src, n, c, margin_mm, title, footer_mm) for n, c in pairs]
    else:
        pdfs = docx_to_pdfs([personalise_docx(src, n, c, title) for n, c in pairs])
    key = io.StringIO()
    w = csv.writer(key)
    w.writerow(["pupil_id", "name"])
    for name, code in pairs:
        w.writerow([code, name])
    return merge_pdfs(pdfs), key.getvalue()


def redact_scan(src: bytes, ext: str, band_mm: float = 25, dpi: int = 200) -> bytes:
    """Black out the top band_mm of every page, flattened to images so nothing is recoverable."""
    ext = ext.lower().lstrip(".")
    if ext == "pdf":
        doc = pymupdf.open(stream=src, filetype="pdf")
    elif ext in ("png", "jpg", "jpeg", "tif", "tiff"):
        img = pymupdf.open(stream=src, filetype=ext)
        doc = pymupdf.open("pdf", img.convert_to_pdf())
    else:
        raise ValueError("Scan must be a PDF or image")
    out = pymupdf.open()
    for page in doc:
        pix = page.get_pixmap(dpi=dpi)
        rows = min(pix.height, round(band_mm / 25.4 * dpi))
        pix = pymupdf.Pixmap(pix, 0) if pix.alpha else pix
        pix.set_rect(pymupdf.IRect(0, 0, pix.width, rows), (0, 0, 0))
        np_ = out.new_page(width=page.rect.width, height=page.rect.height)
        np_.insert_image(np_.rect, pixmap=pix)
    data = out.tobytes(garbage=3, deflate=True)
    out.close()
    doc.close()
    return data


def merge_pdfs(pdfs: list[bytes]) -> bytes:
    """Concatenate PDFs, in order, into one."""
    out = pymupdf.open()
    for data in pdfs:
        with pymupdf.open(stream=data, filetype="pdf") as d:
            out.insert_pdf(d)
    merged = out.tobytes(garbage=3, deflate=True)
    out.close()
    return merged
